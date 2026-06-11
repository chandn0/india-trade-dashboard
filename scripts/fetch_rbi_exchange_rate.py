#!/usr/bin/env python3
import csv
import json
import os
import re
import ssl
import tempfile
import urllib.request

from openpyxl import load_workbook

OUT_JSON = os.path.join('data', 'india_trade_inr_usd_fy.json')
OUT_CSV = os.path.join('data', 'india_trade_inr_usd_fy.csv')
SOURCE_URL = 'https://rbidocs.rbi.org.in/rdocs/Publications/DOCs/139T_13092024245FFE1BB8CB45C3A51183FB6ADA6DC8.XLSX'
SOURCE_NOTE = 'RBI Handbook of Statistics on the Indian Economy, Table 139 (Financial Year - Annual Average and End-year Rates).'


def expand_fy(value: str) -> str:
    match = re.fullmatch(r'(\d{4})-(\d{2})', value)
    if not match:
      return value
    start = int(match.group(1))
    end = int(match.group(2))
    century = (start // 100) * 100
    full_end = century + end
    if full_end <= start:
        full_end += 100
    return f'{start}-{full_end}'


def main():
    ssl_ctx = ssl._create_unverified_context()
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        with urllib.request.urlopen(SOURCE_URL, context=ssl_ctx, timeout=60) as response:
            tmp.write(response.read())
        xlsx_path = tmp.name

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    for row in ws.iter_rows(values_only=True):
        fy = row[1] if len(row) > 1 else None
        if not isinstance(fy, str):
            continue
        if not re.fullmatch(r'\d{4}-\d{2}', fy):
            continue
        rows.append({
            'financial_year': expand_fy(fy),
            'inr_per_usd_avg': float(row[4]) if row[4] is not None else None,
            'inr_per_usd_end': float(row[5]) if row[5] is not None else None,
        })

    payload = {
        'source_url': SOURCE_URL,
        'source_note': SOURCE_NOTE,
        'rows': rows,
    }

    os.makedirs('data', exist_ok=True)
    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(payload, f, indent=2)
        f.write('\n')

    with open(OUT_CSV, 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['financial_year', 'inr_per_usd_avg', 'inr_per_usd_end'])
        for row in rows:
            writer.writerow([row['financial_year'], row['inr_per_usd_avg'], row['inr_per_usd_end']])

    print(f'Wrote {len(rows)} exchange-rate years to {OUT_JSON}')


if __name__ == '__main__':
    main()
